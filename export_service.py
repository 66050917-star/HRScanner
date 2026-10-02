"""ส่งออกตารางฐานข้อมูลเป็น CSV (เปิดใน Excel ได้ ภาษาไทยไม่เพี้ยน) และ Excel (.xlsx)"""
import csv, io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from db_service import HEADERS

def _safe(v):
    """กันสูตรแฝงใน Excel (ข้อความจาก Resume ที่ขึ้นต้นด้วย = + - @)"""
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v

def to_csv(rows) -> bytes:
    buf = io.StringIO(); w = csv.writer(buf); w.writerow(HEADERS)
    for r in rows: w.writerow([_safe(r[h]) if r[h] is not None else "" for h in HEADERS])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM ให้ Excel อ่านภาษาไทยถูก

def to_xlsx(rows) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "ผู้สมัคร"; ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4F46E5"); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows: ws.append([_safe(r[h]) for h in HEADERS])
    widths = [13, 26, 26, 14, 11, 12, 14, 13, 18, 16, 36, 36, 20]
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 32
    out = io.BytesIO(); wb.save(out); return out.getvalue()
